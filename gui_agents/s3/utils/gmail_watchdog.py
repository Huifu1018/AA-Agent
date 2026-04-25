import base64
import csv
import io
import json
import mimetypes
import os
import re
import secrets
import ssl
import threading
import time
import urllib.parse
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import getaddresses, parseaddr, parsedate_to_datetime
from pathlib import Path
from typing import Optional

import anthropic
import certifi
try:
    import docx
except Exception:
    docx = None
try:
    import openpyxl
except Exception:
    openpyxl = None
try:
    import pypdf
except Exception:
    pypdf = None


WATCHDOG_DIR = Path(__file__).resolve().parents[3] / "logs" / "gmail_watchdog"
WATCHDOG_DIR.mkdir(parents=True, exist_ok=True)
DEFAULT_LOG_PATH = WATCHDOG_DIR / "gmail_messages.md"
DEFAULT_STATE_PATH = WATCHDOG_DIR / "gmail_watchdog_state.json"
DEFAULT_CATALOG_BUILD_STATE_PATH = WATCHDOG_DIR / "attachment_catalog_build_state.json"
DEFAULT_TOKEN_PATH = WATCHDOG_DIR / "token.json"
DEFAULT_CREDENTIALS_PATH = WATCHDOG_DIR / "credentials.json"
DEFAULT_AUTH_STATE_PATH = WATCHDOG_DIR / "oauth_state.json"
DEFAULT_ATTACHMENT_CATALOG_PATH = WATCHDOG_DIR / "attachment_catalog.md"
GMAIL_SCOPES = [
    "https://www.googleapis.com/auth/gmail.readonly",
    "https://www.googleapis.com/auth/gmail.compose",
    "https://www.googleapis.com/auth/gmail.modify",
]
SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where())


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _clean_text(text: str) -> str:
    return " ".join((text or "").replace("\x0c", " ").split()).strip()


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _supported_catalog_suffixes() -> set[str]:
    return {
        ".txt",
        ".md",
        ".csv",
        ".pdf",
        ".docx",
        ".xlsx",
    }


def _attachment_catalog_path() -> Path:
    raw = os.getenv("GMAIL_ATTACHMENT_CATALOG_PATH", "").strip()
    return Path(raw).expanduser().resolve() if raw else DEFAULT_ATTACHMENT_CATALOG_PATH


def _ensure_attachment_catalog_template() -> None:
    path = _attachment_catalog_path()
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        """# Gmail 附件知识库

在这里预先登记允许自动作为 Gmail 回复附件发送的本地文件。

使用规则：
- 一条文件一段，按下面的模板填写
- `路径` 必须是文件的绝对路径
- 文件必须落在 Gmail 附件白名单文件夹内，否则不会被自动附加
- `摘要`、`标签`、`关键词` 越清楚，自动选附件越稳

## 文件: 示例客户案例集.pdf
- 路径: /Users/yourname/Documents/materials/示例客户案例集.pdf
- 类型: pdf
- 标签: 客户案例, 智能化转型, OpenClaw
- 摘要: 包含多个客户交流、售前介绍可直接使用的落地案例，适合“请提供案例材料/客户案例/项目介绍”这类邮件。
- 适用场景: 客户交流, 售前介绍, 材料补充
- 禁止场景: 财务数据, 合同, 发票
- 关键词: 案例, 客户案例, 落地案例, 介绍材料, OpenClaw

""",
        encoding="utf-8",
    )


_ensure_attachment_catalog_template()


def _parse_expiry(expiry: str) -> Optional[datetime]:
    if not expiry:
        return None
    value = expiry.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(value).astimezone(timezone.utc)
    except ValueError:
        return None


def _load_oauth_material() -> tuple[dict, Path]:
    token_path = Path(os.getenv("GMAIL_TOKEN_FILE", str(DEFAULT_TOKEN_PATH))).expanduser().resolve()
    if not token_path.exists():
        raise FileNotFoundError(
            f"Gmail token file not found: {token_path}. "
            f"Put an OAuth token JSON there or set GMAIL_TOKEN_FILE."
        )

    token_info = _load_json(token_path)
    if token_info.get("installed") or token_info.get("web"):
        raise RuntimeError(
            "Gmail token file points to OAuth client credentials, not an authorized token. "
            "Use a token.json that contains token/refresh_token."
        )

    credentials_path = Path(
        os.getenv("GMAIL_CREDENTIALS_FILE", str(DEFAULT_CREDENTIALS_PATH))
    ).expanduser().resolve()
    if credentials_path.exists():
        credentials = _load_json(credentials_path)
        block = credentials.get("installed") or credentials.get("web") or {}
        token_info.setdefault("client_id", block.get("client_id"))
        token_info.setdefault("client_secret", block.get("client_secret"))
        token_info.setdefault("token_uri", block.get("token_uri"))

    token_info.setdefault("token_uri", "https://oauth2.googleapis.com/token")
    return token_info, token_path


def _load_client_credentials() -> dict:
    credentials_path = Path(
        os.getenv("GMAIL_CREDENTIALS_FILE", str(DEFAULT_CREDENTIALS_PATH))
    ).expanduser().resolve()
    if credentials_path.exists():
        credentials = _load_json(credentials_path)
        block = credentials.get("installed") or credentials.get("web") or {}
        if block.get("client_id") and block.get("client_secret"):
            return {
                "client_id": block.get("client_id"),
                "client_secret": block.get("client_secret"),
                "auth_uri": block.get("auth_uri", "https://accounts.google.com/o/oauth2/v2/auth"),
                "token_uri": block.get("token_uri", "https://oauth2.googleapis.com/token"),
                "redirect_uris": block.get("redirect_uris", []),
            }
    client_id = os.getenv("GMAIL_CLIENT_ID", "")
    client_secret = os.getenv("GMAIL_CLIENT_SECRET", "")
    if client_id and client_secret:
        return {
            "client_id": client_id,
            "client_secret": client_secret,
            "auth_uri": "https://accounts.google.com/o/oauth2/v2/auth",
            "token_uri": "https://oauth2.googleapis.com/token",
            "redirect_uris": [],
        }
    raise FileNotFoundError(
        f"Gmail OAuth client credentials not found. Put credentials.json at {credentials_path} "
        "or set GMAIL_CLIENT_ID / GMAIL_CLIENT_SECRET."
    )


def _gmail_redirect_uri() -> str:
    return os.getenv("GMAIL_REDIRECT_URI", "http://127.0.0.1:8787/api/gmail-watchdog/oauth/callback")


def _oauth_state_payload() -> dict:
    if DEFAULT_AUTH_STATE_PATH.exists():
        try:
            return _load_json(DEFAULT_AUTH_STATE_PATH)
        except Exception:
            return {}
    return {}


def _save_oauth_state(payload: dict) -> None:
    _write_json(DEFAULT_AUTH_STATE_PATH, payload)


def _oauth_states() -> list[dict]:
    payload = _oauth_state_payload()
    states = payload.get("states")
    if isinstance(states, list):
        return [item for item in states if isinstance(item, dict) and item.get("state")]
    if payload.get("state"):
        return [payload]
    return []


def _remember_oauth_state(state: str) -> None:
    now_ts = int(time.time())
    existing = _oauth_states()
    existing.append(
        {
            "state": state,
            "created_at": _now(),
            "created_at_ts": now_ts,
        }
    )
    deduped: list[dict] = []
    seen: set[str] = set()
    for item in sorted(existing, key=lambda value: int(value.get("created_at_ts", 0)), reverse=True):
        current = str(item.get("state", "")).strip()
        if not current or current in seen:
            continue
        seen.add(current)
        deduped.append(item)
        if len(deduped) >= 8:
            break
    _save_oauth_state({"states": list(reversed(deduped))})


def _consume_oauth_state(state: str, max_age_seconds: int = 1800) -> bool:
    now_ts = int(time.time())
    kept: list[dict] = []
    matched = False
    for item in _oauth_states():
        current = str(item.get("state", "")).strip()
        created_at_ts = int(item.get("created_at_ts", now_ts))
        if current and now_ts - created_at_ts <= max_age_seconds and current == state and not matched:
            matched = True
            continue
        if current and now_ts - created_at_ts <= max_age_seconds:
            kept.append(item)
    _save_oauth_state({"states": kept} if kept else {})
    return matched


def _refresh_access_token(token_info: dict, token_path: Path) -> dict:
    refresh_token = token_info.get("refresh_token")
    client_id = token_info.get("client_id")
    client_secret = token_info.get("client_secret")
    token_uri = token_info.get("token_uri")
    if not (refresh_token and client_id and client_secret and token_uri):
        raise RuntimeError(
            "Gmail OAuth token is missing refresh_token/client_id/client_secret/token_uri."
        )

    payload = urllib.parse.urlencode(
        {
            "client_id": client_id,
            "client_secret": client_secret,
            "refresh_token": refresh_token,
            "grant_type": "refresh_token",
        }
    ).encode("utf-8")
    request = urllib.request.Request(
        token_uri,
        data=payload,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=20, context=SSL_CONTEXT) as response:
        refreshed = json.loads(response.read().decode("utf-8"))

    access_token = refreshed.get("access_token")
    if not access_token:
        raise RuntimeError(f"Failed to refresh Gmail token: {refreshed}")

    token_info["token"] = access_token
    expires_in = int(refreshed.get("expires_in", 3600))
    token_info["expiry"] = (_utc_now() + timedelta(seconds=expires_in - 30)).isoformat()
    _write_json(token_path, token_info)
    return token_info


def gmail_oauth_status() -> dict:
    token_path = Path(os.getenv("GMAIL_TOKEN_FILE", str(DEFAULT_TOKEN_PATH))).expanduser().resolve()
    credentials_path = Path(
        os.getenv("GMAIL_CREDENTIALS_FILE", str(DEFAULT_CREDENTIALS_PATH))
    ).expanduser().resolve()
    token_exists = token_path.exists()
    credentials_exists = credentials_path.exists() or bool(
        os.getenv("GMAIL_CLIENT_ID") and os.getenv("GMAIL_CLIENT_SECRET")
    )
    expiry = ""
    if token_exists:
        try:
            token_info = _load_json(token_path)
            expiry = token_info.get("expiry", "")
        except Exception:
            expiry = ""
    return {
        "token_path": str(token_path),
        "credentials_path": str(credentials_path),
        "token_exists": token_exists,
        "credentials_exists": credentials_exists,
        "redirect_uri": _gmail_redirect_uri(),
        "expiry": expiry,
    }


def gmail_attachment_catalog_status() -> dict:
    path = _attachment_catalog_path()
    return {
        "catalog_path": str(path),
        "catalog_exists": path.exists(),
    }


def gmail_attachment_catalog_preview(limit: int = 12) -> list[dict]:
    return _parse_attachment_catalog()[:limit]


def start_gmail_oauth_flow() -> dict:
    client = _load_client_credentials()
    state = secrets.token_urlsafe(24)
    _remember_oauth_state(state)
    query = urllib.parse.urlencode(
        {
            "client_id": client["client_id"],
            "redirect_uri": _gmail_redirect_uri(),
            "response_type": "code",
            "scope": " ".join(GMAIL_SCOPES),
            "access_type": "offline",
            "prompt": "consent",
            "include_granted_scopes": "true",
            "state": state,
        }
    )
    return {
        "auth_url": f"{client['auth_uri']}?{query}",
        "state": state,
        **gmail_oauth_status(),
    }


def finish_gmail_oauth_flow(code: str, state: str) -> dict:
    if not code:
        raise RuntimeError("Missing Gmail OAuth code.")
    if not state or not _consume_oauth_state(state):
        raise RuntimeError("Invalid Gmail OAuth state.")

    client = _load_client_credentials()
    payload = urllib.parse.urlencode(
        {
            "code": code,
            "client_id": client["client_id"],
            "client_secret": client["client_secret"],
            "redirect_uri": _gmail_redirect_uri(),
            "grant_type": "authorization_code",
        }
    ).encode("utf-8")
    request = urllib.request.Request(
        client["token_uri"],
        data=payload,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=30, context=SSL_CONTEXT) as response:
        token_response = json.loads(response.read().decode("utf-8"))

    if not token_response.get("access_token"):
        raise RuntimeError(f"Gmail OAuth exchange failed: {token_response}")

    token_path = Path(os.getenv("GMAIL_TOKEN_FILE", str(DEFAULT_TOKEN_PATH))).expanduser().resolve()
    token_payload = {
        "token": token_response.get("access_token"),
        "refresh_token": token_response.get("refresh_token"),
        "scope": token_response.get("scope", " ".join(GMAIL_SCOPES)),
        "token_type": token_response.get("token_type", "Bearer"),
        "expiry": (_utc_now() + timedelta(seconds=int(token_response.get("expires_in", 3600) - 30))).isoformat(),
        "client_id": client["client_id"],
        "client_secret": client["client_secret"],
        "token_uri": client["token_uri"],
    }
    if not token_payload["refresh_token"] and token_path.exists():
        try:
            previous = _load_json(token_path)
            token_payload["refresh_token"] = previous.get("refresh_token", "")
        except Exception:
            pass
    _write_json(token_path, token_payload)
    _save_oauth_state({})
    return gmail_oauth_status()


def _authorized_token() -> tuple[str, dict]:
    token_info, token_path = _load_oauth_material()
    token = token_info.get("token")
    expiry = _parse_expiry(token_info.get("expiry", ""))
    if not token or (expiry and expiry <= _utc_now()):
        token_info = _refresh_access_token(token_info, token_path)
        token = token_info.get("token")
    if not token:
        raise RuntimeError("No Gmail access token available.")
    return token, token_info


def _gmail_api_get(path: str, query: Optional[dict] = None) -> dict:
    token, _ = _authorized_token()
    url = "https://gmail.googleapis.com/gmail/v1/" + path.lstrip("/")
    if query:
        url += "?" + urllib.parse.urlencode(query)
    request = urllib.request.Request(
        url,
        headers={"Authorization": f"Bearer {token}"},
        method="GET",
    )
    with urllib.request.urlopen(request, timeout=30, context=SSL_CONTEXT) as response:
        return json.loads(response.read().decode("utf-8"))


def _gmail_api_post(path: str, payload: dict) -> dict:
    token, _ = _authorized_token()
    url = "https://gmail.googleapis.com/gmail/v1/" + path.lstrip("/")
    request = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=30, context=SSL_CONTEXT) as response:
        return json.loads(response.read().decode("utf-8"))


def _gmail_api_get_json(path: str) -> dict:
    token, _ = _authorized_token()
    url = "https://gmail.googleapis.com/gmail/v1/" + path.lstrip("/")
    request = urllib.request.Request(
        url,
        headers={"Authorization": f"Bearer {token}"},
        method="GET",
    )
    with urllib.request.urlopen(request, timeout=30, context=SSL_CONTEXT) as response:
        return json.loads(response.read().decode("utf-8"))


def _decode_body(data: str) -> str:
    if not data:
        return ""
    padding = "=" * (-len(data) % 4)
    raw = base64.urlsafe_b64decode((data + padding).encode("utf-8"))
    return raw.decode("utf-8", errors="replace")


def _extract_plain_text(payload: dict) -> str:
    if not payload:
        return ""
    mime_type = payload.get("mimeType", "")
    body_data = payload.get("body", {}).get("data", "")
    if mime_type.startswith("text/plain") and body_data:
        return _clean_text(_decode_body(body_data))
    parts = payload.get("parts") or []
    for part in parts:
        text = _extract_plain_text(part)
        if text:
            return text
    if body_data:
        return _clean_text(_decode_body(body_data))
    return ""


def _extract_latest_email_body(text: str) -> str:
    cleaned = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not cleaned:
        return ""

    lines = [line.rstrip() for line in cleaned.split("\n")]
    cut_patterns = [
        re.compile(r"^发件人[:：]"),
        re.compile(r"^寄件者[:：]"),
        re.compile(r"^发送时间[:：]"),
        re.compile(r"^日期[:：]"),
        re.compile(r"^收件人[:：]"),
        re.compile(r"^抄送[:：]"),
        re.compile(r"^主题[:：]"),
        re.compile(r"^From:\s", re.IGNORECASE),
        re.compile(r"^Sent:\s", re.IGNORECASE),
        re.compile(r"^Date:\s", re.IGNORECASE),
        re.compile(r"^To:\s", re.IGNORECASE),
        re.compile(r"^Cc:\s", re.IGNORECASE),
        re.compile(r"^Subject:\s", re.IGNORECASE),
        re.compile(r"^On .+wrote:\s*$", re.IGNORECASE),
        re.compile(r"^-{2,}\s*Original Message\s*-{2,}$", re.IGNORECASE),
    ]

    cutoff = len(lines)
    for idx, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue
        if any(pattern.match(stripped) for pattern in cut_patterns):
            cutoff = idx
            break

    latest = "\n".join(lines[:cutoff]).strip() or cleaned
    return latest


def _iter_attachment_parts(payload: dict) -> list[dict]:
    parts = payload.get("parts") or []
    found: list[dict] = []
    for part in parts:
        filename = (part.get("filename") or "").strip()
        body = part.get("body", {}) or {}
        if filename and (body.get("attachmentId") or body.get("data")):
            found.append(part)
        found.extend(_iter_attachment_parts(part))
    return found


def _fetch_attachment_bytes(message_id: str, part: dict) -> bytes:
    body = part.get("body", {}) or {}
    data = body.get("data")
    if not data and body.get("attachmentId"):
        payload = _gmail_api_get_json(
            f"users/me/messages/{message_id}/attachments/{body['attachmentId']}"
        )
        data = payload.get("data", "")
    if not data:
        return b""
    padding = "=" * (-len(data) % 4)
    return base64.urlsafe_b64decode((data + padding).encode("utf-8"))


def _truncate_text(text: str, limit: int = 4000) -> str:
    stripped = (text or "").strip()
    if len(stripped) <= limit:
        return stripped
    return stripped[:limit] + "\n...[truncated]..."


def _extract_attachment_text(filename: str, mime_type: str, raw: bytes) -> str:
    name = (filename or "").lower()
    mime = (mime_type or "").lower()
    if not raw:
        return ""
    try:
        if name.endswith(".txt") or name.endswith(".md") or mime.startswith("text/plain") or mime == "text/markdown":
            return raw.decode("utf-8", errors="replace")
        if name.endswith(".csv") or mime.startswith("text/csv"):
            return raw.decode("utf-8", errors="replace")
        if name.endswith(".pdf") or mime == "application/pdf":
            if pypdf is None:
                return "[当前环境未安装 pypdf，暂不解析 PDF 文本]"
            reader = pypdf.PdfReader(io.BytesIO(raw))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if name.endswith(".docx") or mime.endswith("wordprocessingml.document"):
            if docx is None:
                return "[当前环境未安装 python-docx，暂不解析 DOCX 文本]"
            doc = docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if name.endswith(".xlsx") or mime.endswith("spreadsheetml.sheet"):
            if openpyxl is None:
                return "[当前环境未安装 openpyxl，暂不解析 XLSX 文本]"
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            rows: list[str] = []
            for ws in wb.worksheets[:3]:
                rows.append(f"[Sheet] {ws.title}")
                for row in ws.iter_rows(max_row=20, values_only=True):
                    values = [str(v).strip() for v in row if v not in (None, "")]
                    if values:
                        rows.append(" | ".join(values))
            return "\n".join(rows)
    except Exception as exc:
        return f"[附件解析失败: {type(exc).__name__}]"
    return ""


def _extract_local_file_text(path: Path) -> str:
    resolved = path.expanduser().resolve()
    mime_type, _ = mimetypes.guess_type(str(resolved))
    raw = resolved.read_bytes()
    return _extract_attachment_text(resolved.name, mime_type or "", raw)


def _collect_catalog_terms(file_path: Path, text: str, limit: int = 8) -> list[str]:
    source = f"{file_path.stem}\n{text[:1200]}"
    tokens = re.findall(r"[\u4e00-\u9fff]{2,}|[A-Za-z][A-Za-z0-9_\-]{2,}", source)
    blocked = {
        "的", "我们", "你们", "他们", "以及", "这个", "那个", "进行", "需要", "相关", "内容", "文件",
        "materials", "report", "analysis", "attachment", "document", "sheet", "page",
    }
    result: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        normalized = token.strip()
        lowered = normalized.lower()
        if not normalized or lowered in blocked or lowered in seen:
            continue
        seen.add(lowered)
        result.append(normalized)
        if len(result) >= limit:
            break
    return result


def _infer_use_cases(file_path: Path, text: str, keywords: list[str]) -> tuple[str, str]:
    haystack = f"{file_path.name}\n{text[:1600]}\n{' '.join(keywords)}"
    rules = [
        (("案例", "客户", "落地", "方案"), ("客户交流, 售前介绍, 材料补充", "合同, 发票, 内部保密内容")),
        (("年报", "季报", "财务", "经营"), ("经营分析, 财务解读, 管理汇报", "未公开披露场景, 合同回复")),
        (("ppt", "演讲", "汇报", "介绍"), ("会议交流, 汇报材料, 客户介绍", "合同, 发票")),
        (("简历", "履历", "候选人"), ("候选人沟通, 背景介绍", "客户材料, 合同")),
    ]
    lower_haystack = haystack.lower()
    for needles, values in rules:
        if any(needle.lower() in lower_haystack for needle in needles):
            return values
    return ("资料补充, 邮件回复附件", "未确认内容准确性前不要外发")


def _build_catalog_entry(file_path: Path, text: str) -> dict:
    cleaned = _clean_text(text)
    if not cleaned:
        return {
            "name": file_path.name,
            "path": str(file_path),
            "type": file_path.suffix.lstrip(".").lower() or "unknown",
            "tags": [],
            "summary": f"文件 {file_path.name} 当前未提取到可用正文，可人工补充摘要。",
            "use_cases": "材料补充",
            "avoid_cases": "未确认内容准确性前不要外发",
            "keywords": [],
            "excerpt": "",
        }

    catalog_summary = _summarize_catalog_entry(file_path, cleaned)
    keywords = _split_meta_items(catalog_summary.get("keywords", "")) or _collect_catalog_terms(file_path, cleaned)
    tags = _split_meta_items(catalog_summary.get("tags", ""), limit=6) or keywords[:4]
    use_cases = catalog_summary.get("use_cases", "").strip()
    avoid_cases = catalog_summary.get("avoid_cases", "").strip()
    if not use_cases or not avoid_cases:
        inferred_use_cases, inferred_avoid_cases = _infer_use_cases(file_path, cleaned, keywords)
        use_cases = use_cases or inferred_use_cases
        avoid_cases = avoid_cases or inferred_avoid_cases

    return {
        "name": file_path.name,
        "path": str(file_path),
        "type": file_path.suffix.lstrip(".").lower() or "unknown",
        "tags": tags,
        "summary": catalog_summary.get("summary", "").strip() or _heuristic_summary(file_path, cleaned),
        "use_cases": use_cases,
        "avoid_cases": avoid_cases,
        "keywords": keywords,
        "excerpt": catalog_summary.get("excerpt", "").strip() or _truncate_text(cleaned, limit=1200),
    }


def _render_attachment_catalog(entries: list[dict]) -> str:
    lines = [
        "# Gmail 附件知识库",
        "",
        f"> 自动初始化时间：{_now()}",
        "",
        "以下条目由 AA-CUA 根据白名单文件夹自动整理，可在此基础上人工补充或修正。",
        "",
    ]
    for entry in entries:
        lines.extend(
            [
                f"## 文件: {entry['name']}",
                f"- 路径: {entry['path']}",
                f"- 类型: {entry['type']}",
                f"- 标签: {', '.join(entry['tags']) if entry['tags'] else '待补充'}",
                f"- 摘要: {entry['summary']}",
                f"- 适用场景: {entry['use_cases']}",
                f"- 禁止场景: {entry['avoid_cases']}",
                f"- 关键词: {', '.join(entry['keywords']) if entry['keywords'] else '待补充'}",
                "- 内容摘录:",
                "",
                "```",
                (entry.get("excerpt", "") or "").strip(),
                "```",
                "",
            ]
        )
    if len(lines) <= 6:
        lines.extend(
            [
                "## 文件: 暂无",
                "- 路径: ",
                "- 类型: ",
                "- 标签: 待补充",
                "- 摘要: 当前白名单文件夹里还没有生成任何可用条目。",
                "- 适用场景: 待补充",
                "- 禁止场景: 待补充",
                "- 关键词: 待补充",
                "",
            ]
        )
    return "\n".join(lines).rstrip() + "\n"


def _extract_attachment_summaries(message_id: str, payload: dict) -> list[dict]:
    summaries: list[dict] = []
    for part in _iter_attachment_parts(payload):
        filename = (part.get("filename") or "").strip() or "unnamed"
        mime_type = part.get("mimeType", "")
        raw = _fetch_attachment_bytes(message_id, part)
        text = _extract_attachment_text(filename, mime_type, raw)
        summaries.append(
            {
                "filename": filename,
                "mime_type": mime_type,
                "text": _truncate_text(_clean_text(text), limit=2500) if text else "",
            }
        )
    return summaries


def _format_attachment_context(attachments: list[dict]) -> str:
    chunks: list[str] = []
    for item in attachments or []:
        filename = item.get("filename", "unnamed")
        text = (item.get("text") or "").strip()
        if text:
            chunks.append(f"[附件] {filename}\n{text}")
        else:
            chunks.append(f"[附件] {filename}\n[无可提取文本内容]")
    return "\n\n".join(chunks).strip()


def _message_headers(payload: dict) -> dict:
    result = {}
    for header in payload.get("headers") or []:
        name = (header.get("name") or "").lower()
        if name in {"from", "subject", "date", "message-id", "references", "reply-to", "to", "cc"}:
            result[name] = header.get("value", "")
    return result


def _require_compose_scope() -> None:
    token_info, _ = _load_oauth_material()
    scope_text = token_info.get("scope", "")
    scopes = set(scope_text.split()) if scope_text else set()
    if (
        "https://www.googleapis.com/auth/gmail.compose" not in scopes
        and "https://www.googleapis.com/auth/gmail.send" not in scopes
    ):
        raise RuntimeError(
            "当前 Gmail 授权没有发送权限。请在验证台重新点击一次“连接 Gmail”完成升级授权。"
        )


def _require_modify_scope() -> None:
    token_info, _ = _load_oauth_material()
    scope_text = token_info.get("scope", "")
    scopes = set(scope_text.split()) if scope_text else set()
    if "https://www.googleapis.com/auth/gmail.modify" not in scopes:
        raise RuntimeError(
            "当前 Gmail 授权没有 modify 权限。请在验证台重新点击一次“连接 Gmail”完成升级授权。"
        )


def find_gmail_message_by_subject(subject: str, unread_only: bool = False) -> dict:
    query = f'subject:\"{subject}\"'
    if unread_only:
        query = f"is:unread {query}"
    payload = _gmail_api_get(
        "users/me/messages",
        {"q": query, "maxResults": 1},
    )
    messages = payload.get("messages", []) or []
    if not messages:
        raise FileNotFoundError(f"没有找到主题包含“{subject}”的邮件。")
    return _fetch_message_detail_by_id(messages[0]["id"])


def _fetch_message_detail_by_id(message_id: str) -> dict:
    payload = _gmail_api_get(
        f"users/me/messages/{message_id}",
        {"format": "full"},
    )
    headers = _message_headers(payload.get("payload", {}))
    raw_body = _extract_plain_text(payload.get("payload", {})) or _clean_text(payload.get("snippet", ""))
    body = _extract_latest_email_body(raw_body)
    attachments = _extract_attachment_summaries(message_id, payload.get("payload", {}))
    date_value = headers.get("date", "")
    try:
        date_text = parsedate_to_datetime(date_value).astimezone().strftime("%Y-%m-%d %H:%M:%S") if date_value else ""
    except Exception:
        date_text = date_value
    return {
        "id": payload.get("id", message_id),
        "threadId": payload.get("threadId", ""),
        "from": headers.get("from", ""),
        "subject": headers.get("subject", ""),
        "date": date_text,
        "snippet": _clean_text(payload.get("snippet", "")),
        "body": body,
        "attachments": attachments,
        "attachment_context": _format_attachment_context(attachments),
        "message_id_header": headers.get("message-id", ""),
        "references": headers.get("references", ""),
        "reply_to": headers.get("reply-to", ""),
        "to": headers.get("to", ""),
        "cc": headers.get("cc", ""),
    }


def _normalize_email_list(header_value: str) -> list[str]:
    emails: list[str] = []
    for _, email in getaddresses([header_value or ""]):
        normalized = (email or "").strip().lower()
        if normalized and normalized not in emails:
            emails.append(normalized)
    return emails


def _get_authenticated_email() -> str:
    profile = _gmail_api_get("users/me/profile")
    return (profile.get("emailAddress") or "").strip().lower()


def _format_quoted_reply_block(message: dict) -> str:
    sender = message.get("from", "").strip() or "未知发件人"
    date_text = message.get("date", "").strip() or "未知时间"
    subject = message.get("subject", "").strip() or "无主题"
    body = (message.get("body", "") or message.get("snippet", "") or "").strip()
    if not body:
        body = "(原邮件正文为空)"
    quoted_lines = "\n".join(f"> {line}" if line else ">" for line in body.splitlines())
    return (
        "\n\n"
        f"在 {date_text}，{sender} 写道：\n"
        f"> Subject: {subject}\n"
        f"{quoted_lines}"
    )


def _sanitize_reply_body(text: str) -> str:
    cleaned = (text or "").strip()
    if not cleaned:
        return ""
    replacements = {
        "**": "",
        "__": "",
        "### ": "",
        "## ": "",
        "# ": "",
    }
    for old, new in replacements.items():
        cleaned = cleaned.replace(old, new)
    lines = []
    for raw_line in cleaned.splitlines():
        line = raw_line.rstrip()
        stripped = line.lstrip()
        if stripped.startswith(("- ", "* ")):
            line = stripped[2:].strip()
        lines.append(line)
    cleaned = "\n".join(lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _gmail_local_attachment_roots() -> list[Path]:
    raw = os.getenv("GMAIL_REPLY_ATTACHMENT_DIRS", "").strip()
    if not raw:
        return []
    roots: list[Path] = []
    for chunk in raw.split(os.pathsep):
        candidate = Path(chunk.strip()).expanduser()
        if candidate.exists() and candidate.is_dir():
            roots.append(candidate.resolve())
    deduped: list[Path] = []
    seen: set[Path] = set()
    for root in roots:
        if root not in seen:
            deduped.append(root)
            seen.add(root)
    return deduped


def _path_within_roots(path: Path, roots: list[Path]) -> bool:
    resolved = path.expanduser().resolve()
    for root in roots:
        try:
            resolved.relative_to(root)
            return True
        except Exception:
            continue
    return False


def _parse_attachment_catalog() -> list[dict]:
    path = _attachment_catalog_path()
    if not path.exists():
        return []
    text = path.read_text(encoding="utf-8")
    sections = re.split(r"(?m)^##\s+", text)
    entries: list[dict] = []
    for section in sections:
        chunk = section.strip()
        if not chunk:
            continue
        lines = chunk.splitlines()
        heading = lines[0].strip()
        body_lines = lines[1:]
        fields: dict[str, str] = {}
        excerpt_lines: list[str] = []
        collecting_excerpt = False
        in_fence = False
        for raw_line in body_lines:
            line = raw_line.strip()
            if collecting_excerpt:
                if line == "```":
                    if in_fence:
                        collecting_excerpt = False
                        in_fence = False
                    else:
                        in_fence = True
                    continue
                if in_fence:
                    excerpt_lines.append(raw_line.rstrip())
                continue
            if not line.startswith("- "):
                continue
            payload = line[2:]
            if ":" not in payload:
                continue
            key, value = payload.split(":", 1)
            fields[key.strip()] = value.strip()
            if key.strip() == "内容摘录":
                collecting_excerpt = True
        path_value = fields.get("路径", "").strip()
        if not path_value:
            continue
        try:
            file_path = Path(path_value).expanduser().resolve()
        except Exception:
            continue
        entries.append(
            {
                "heading": heading,
                "declared_name": re.sub(r"^文件:\s*", "", heading).strip(),
                "path": str(file_path),
                "name": file_path.name,
                "suffix": file_path.suffix.lower(),
                "type": fields.get("类型", "").strip(),
                "tags": fields.get("标签", "").strip(),
                "summary": fields.get("摘要", "").strip(),
                "use_cases": fields.get("适用场景", "").strip(),
                "avoid_cases": fields.get("禁止场景", "").strip(),
                "keywords": fields.get("关键词", "").strip(),
                "excerpt": "\n".join(excerpt_lines).strip(),
            }
        )
    return entries


def _catalog_attachment_candidates(limit: int = 80) -> list[dict]:
    roots = _gmail_local_attachment_roots()
    if not roots:
        return []
    allowed_exts = {
        ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt", ".md",
        ".ppt", ".pptx", ".png", ".jpg", ".jpeg", ".zip",
    }
    entries = _parse_attachment_catalog()
    candidates: list[dict] = []
    for entry in entries:
        path = Path(entry["path"])
        if not path.exists() or not path.is_file():
            continue
        if path.suffix.lower() not in allowed_exts:
            continue
        if not _path_within_roots(path, roots):
            continue
        try:
            stat = path.stat()
        except Exception:
            continue
        candidates.append(
            {
                "path": str(path),
                "name": path.name,
                "suffix": path.suffix.lower(),
                "size": stat.st_size,
                "mtime_ts": float(stat.st_mtime),
                "mtime": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
                "root": next((str(root) for root in roots if _path_within_roots(path, [root])), ""),
                "summary": entry.get("summary", ""),
                "tags": entry.get("tags", ""),
                "keywords": entry.get("keywords", ""),
                "use_cases": entry.get("use_cases", ""),
                "avoid_cases": entry.get("avoid_cases", ""),
                "source": "catalog",
            }
        )
    candidates.sort(key=lambda item: (-float(item.get("mtime_ts", 0.0)), item["path"]))
    return candidates[:limit]


def _read_local_file_text(path: Path) -> str:
    suffix = path.suffix.lower()
    mime_type = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
    raw = path.read_bytes()
    return _extract_attachment_text(path.name, mime_type, raw)


def _heuristic_keywords(text: str, path: Path) -> str:
    source = f"{path.stem} {(text or '')}"
    parts = re.findall(r"[\u4e00-\u9fffA-Za-z0-9]{2,20}", source)
    seen: list[str] = []
    for token in parts:
        value = token.strip()
        if not value:
            continue
        lowered = value.lower()
        if lowered in {"users", "documents", "report", "final", "draft"}:
            continue
        if value not in seen:
            seen.append(value)
        if len(seen) >= 8:
            break
    return ", ".join(seen)


def _split_meta_items(value: str, limit: int = 8) -> list[str]:
    parts = re.split(r"[,，、;\n]+", value or "")
    items: list[str] = []
    for part in parts:
        cleaned = part.strip()
        if not cleaned:
            continue
        if cleaned not in items:
            items.append(cleaned)
        if len(items) >= limit:
            break
    return items


def _heuristic_summary(path: Path, cleaned: str) -> str:
    chunks = [item.strip() for item in re.split(r"[\n。！？!?\r]+", cleaned or "") if item.strip()]
    selected: list[str] = []
    for chunk in chunks:
        if len(chunk) < 12:
            continue
        selected.append(chunk)
        if len(selected) >= 3:
            break
    core = "；".join(selected)
    if not core:
        core = f"{path.stem} 的资料文件。"
    summary = (
        f"该材料主要围绕“{path.stem}”展开，核心内容包括：{core}。"
        "适合在邮件中用于介绍材料主题、补充背景信息、提供分析结论或作为附件说明发送。"
    )
    return _truncate_text(summary, limit=260)


def _summarize_catalog_entry(path: Path, text: str) -> dict:
    cleaned = _clean_text(text)
    excerpt = _truncate_text(cleaned, limit=800)
    api_key = os.getenv("KIMI_API_KEY") or os.getenv("ANTHROPIC_API_KEY")
    if not api_key or not cleaned:
        return {
            "summary": _heuristic_summary(path, cleaned),
            "tags": path.suffix.lower().lstrip("."),
            "use_cases": "材料补充",
            "avoid_cases": "",
            "keywords": _heuristic_keywords(cleaned, path),
            "excerpt": excerpt,
        }
    model = os.getenv("KIMI_MODEL", os.getenv("ANTHROPIC_MODEL", "kimi-k2.5"))
    base_url = os.getenv("KIMI_BASE_URL", os.getenv("ANTHROPIC_BASE_URL", "https://api.moonshot.cn/anthropic"))
    client = anthropic.Anthropic(api_key=api_key, base_url=base_url)
    prompt = (
        "你在为 Gmail 自动回复附件系统生成附件知识库条目。\n"
        "请基于下面文件内容，输出 JSON："
        '{"summary":"一句中文摘要","tags":"逗号分隔标签","use_cases":"逗号分隔适用场景","avoid_cases":"逗号分隔禁止场景","keywords":"逗号分隔关键词"}\n'
        "要求：\n"
        "1. summary 用 2-4 句中文完整介绍材料内容，不少于 80 字，尽量说明主题、核心内容、适用场景；\n"
        "2. 不要只截取开头文本，不要直接输出 OCR 噪声；\n"
        "3. tags / use_cases / avoid_cases / keywords 保持简短，使用逗号分隔；\n"
        "4. 输出必须是严格 JSON，不要添加解释。\n\n"
        f"文件名：{path.name}\n"
        f"文件内容：{_truncate_text(cleaned, limit=3500)}"
    )
    try:
        response = client.messages.create(
            model=model,
            max_tokens=320,
            messages=[{"role": "user", "content": prompt}],
        )
        text_resp = "\n".join(getattr(block, "text", "") for block in response.content).strip()
        match = re.search(r"\{.*\}", text_resp, re.S)
        if match:
            payload = json.loads(match.group(0))
            return {
                "summary": str(payload.get("summary", "")).strip() or _heuristic_summary(path, cleaned),
                "tags": str(payload.get("tags", "")).strip(),
                "use_cases": str(payload.get("use_cases", "")).strip(),
                "avoid_cases": str(payload.get("avoid_cases", "")).strip(),
                "keywords": str(payload.get("keywords", "")).strip() or _heuristic_keywords(cleaned, path),
                "excerpt": excerpt,
            }
    except Exception:
        pass
    return {
        "summary": _heuristic_summary(path, cleaned),
        "tags": path.suffix.lower().lstrip("."),
        "use_cases": "材料补充",
        "avoid_cases": "",
        "keywords": _heuristic_keywords(cleaned, path),
        "excerpt": excerpt,
    }


@dataclass
class AttachmentCatalogBuildSnapshot:
    running: bool = False
    processed: int = 0
    total: int = 0
    current_file: str = ""
    last_started_at: str = ""
    last_finished_at: str = ""
    last_error: str = ""
    last_message: str = ""
    catalog_path: str = str(DEFAULT_ATTACHMENT_CATALOG_PATH)


class AttachmentCatalogBuilder:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._thread: Optional[threading.Thread] = None
        self._snapshot = AttachmentCatalogBuildSnapshot(catalog_path=str(_attachment_catalog_path()))

    def _snapshot_dict(self) -> dict:
        return {
            "running": self._snapshot.running,
            "processed": self._snapshot.processed,
            "total": self._snapshot.total,
            "current_file": self._snapshot.current_file,
            "last_started_at": self._snapshot.last_started_at,
            "last_finished_at": self._snapshot.last_finished_at,
            "last_error": self._snapshot.last_error,
            "last_message": self._snapshot.last_message,
            "catalog_path": self._snapshot.catalog_path,
        }

    def status(self) -> dict:
        with self._lock:
            self._snapshot.catalog_path = str(_attachment_catalog_path())
            return self._snapshot_dict()

    def start(self) -> dict:
        with self._lock:
            if self._thread and self._thread.is_alive():
                return self._snapshot_dict()
            self._snapshot = AttachmentCatalogBuildSnapshot(
                running=True,
                processed=0,
                total=0,
                current_file="",
                last_started_at=_now(),
                last_finished_at="",
                last_error="",
                last_message="正在初始化附件知识库...",
                catalog_path=str(_attachment_catalog_path()),
            )
            self._thread = threading.Thread(target=self._build, daemon=True)
            self._thread.start()
            return self._snapshot_dict()

    def _build(self) -> None:
        try:
            roots = _gmail_local_attachment_roots()
            allowed_exts = {
                ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt", ".md",
                ".ppt", ".pptx", ".png", ".jpg", ".jpeg", ".zip",
            }
            files: list[Path] = []
            for root in roots:
                for path in root.rglob("*"):
                    if path.is_file() and path.suffix.lower() in allowed_exts:
                        files.append(path.resolve())
            files = sorted(dict.fromkeys(files))
            with self._lock:
                self._snapshot.total = len(files)
                self._snapshot.last_message = f"共发现 {len(files)} 个候选文件，开始分析。"
            entries: list[str] = [
                "# Gmail 附件知识库",
                "",
                "这份文档由初始化流程生成，后续可以手工补充或修改。",
                "",
            ]
            for idx, path in enumerate(files, start=1):
                with self._lock:
                    self._snapshot.current_file = str(path)
                    self._snapshot.processed = idx - 1
                    self._snapshot.last_message = f"正在分析：{path.name}"
                try:
                    text = _read_local_file_text(path)
                except Exception as exc:
                    text = f"[读取失败: {type(exc).__name__}]"
                summary = _summarize_catalog_entry(path, text)
                entries.extend(
                    [
                        f"## 文件: {path.name}",
                        f"- 路径: {path}",
                        f"- 类型: {path.suffix.lower().lstrip('.') or 'unknown'}",
                        f"- 标签: {summary.get('tags', '')}",
                        f"- 摘要: {summary.get('summary', '')}",
                        f"- 适用场景: {summary.get('use_cases', '')}",
                        f"- 禁止场景: {summary.get('avoid_cases', '')}",
                        f"- 关键词: {summary.get('keywords', '')}",
                        "- 内容摘录:",
                        "",
                        "```",
                        (summary.get("excerpt", "") or "").strip(),
                        "```",
                        "",
                    ]
                )
                with self._lock:
                    self._snapshot.processed = idx
            catalog_path = _attachment_catalog_path()
            catalog_path.parent.mkdir(parents=True, exist_ok=True)
            catalog_path.write_text("\n".join(entries).strip() + "\n", encoding="utf-8")
            with self._lock:
                self._snapshot.running = False
                self._snapshot.current_file = ""
                self._snapshot.last_finished_at = _now()
                self._snapshot.last_message = "附件知识库初始化完成，可以启动 Gmail Watchdog。"
                self._snapshot.catalog_path = str(catalog_path)
        except Exception as exc:
            with self._lock:
                self._snapshot.running = False
                self._snapshot.last_finished_at = _now()
                self._snapshot.last_error = str(exc)
                self._snapshot.last_message = "附件知识库初始化失败。"


CATALOG_BUILDER = AttachmentCatalogBuilder()


def _list_local_attachment_candidates(limit: int = 80) -> list[dict]:
    roots = _gmail_local_attachment_roots()
    if not roots:
        return []
    allowed_exts = {
        ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt", ".md",
        ".ppt", ".pptx", ".png", ".jpg", ".jpeg", ".zip",
    }
    candidates: list[dict] = []
    for root in roots:
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            if path.suffix.lower() not in allowed_exts:
                continue
            try:
                stat = path.stat()
            except Exception:
                continue
            candidates.append(
                {
                    "path": str(path),
                    "name": path.name,
                    "suffix": path.suffix.lower(),
                    "size": stat.st_size,
                    "mtime_ts": float(stat.st_mtime),
                    "mtime": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
                    "root": str(root),
                }
            )
    candidates.sort(key=lambda item: (-float(item.get("mtime_ts", 0.0)), item["path"]))
    return candidates[:limit]


def _select_local_reply_attachments(
    sender: str,
    subject: str,
    content: str,
    candidate_files: list[dict],
) -> dict:
    if not candidate_files:
        return {
            "use_catalog": False,
            "should_attach": False,
            "reason": "No configured local attachment directories.",
            "selected_paths": [],
        }
    api_key = os.getenv("KIMI_API_KEY") or os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        return {
            "use_catalog": False,
            "should_attach": False,
            "reason": "No Kimi/Anthropic API key available for attachment selection.",
            "selected_paths": [],
        }
    model = os.getenv("KIMI_MODEL", os.getenv("ANTHROPIC_MODEL", "kimi-k2.5"))
    base_url = os.getenv("KIMI_BASE_URL", os.getenv("ANTHROPIC_BASE_URL", "https://api.moonshot.cn/anthropic"))
    client = anthropic.Anthropic(api_key=api_key, base_url=base_url)
    file_lines = []
    for idx, item in enumerate(candidate_files, start=1):
        details: list[str] = [
            f"{idx}. {item['name']}",
            f"path={item['path']}",
            f"size={item['size']}",
            f"modified={item['mtime']}",
        ]
        if item.get("summary"):
            details.append(f"摘要={item['summary']}")
        if item.get("tags"):
            details.append(f"标签={item['tags']}")
        if item.get("keywords"):
            details.append(f"关键词={item['keywords']}")
        if item.get("use_cases"):
            details.append(f"适用场景={item['use_cases']}")
        if item.get("avoid_cases"):
            details.append(f"禁止场景={item['avoid_cases']}")
        if item.get("excerpt"):
            details.append(f"内容摘录={_truncate_text(str(item['excerpt']), limit=220)}")
        file_lines.append(
            " | ".join(details)
        )
    has_catalog = any((item.get("source") == "catalog") for item in candidate_files)
    catalog_rule = (
        "这些候选文件来自预先维护的附件知识库文档，知识库里的摘要、标签、关键词和适用场景应作为主要选择依据。\n"
        if has_catalog
        else ""
    )
    prompt = (
        "你在做 Gmail 自动回复的本地附件选择。\n"
        "请判断是否需要参考本地资料知识库中的文档内容来帮助回复，以及是否需要把这些文件作为附件一起发出去。\n"
        "规则：\n"
        "1. 如果邮件问题明显涉及本地资料目录里的材料内容、报告结论、案例说明、分析结果、PPT/文档内容，请选出相关文件作为知识库参考。\n"
        "2. 只有在邮件明确索要本地材料、方案、PPT、文档、报告、图片、附件时，才允许 should_attach=true。\n"
        "3. 如果只是需要借助本地资料内容来回答，但对方没有明确要附件，可以 use_catalog=true 且 should_attach=false。\n"
        "4. 如果不确定，就不要选。\n\n"
        f"发件人：{sender}\n"
        f"主题：{subject}\n"
        f"邮件内容：{content}\n\n"
        f"{catalog_rule}"
        "可选本地文件如下：\n"
        + "\n".join(file_lines)
        + "\n\n请只输出 JSON，格式如下：\n"
        '{"use_catalog": true/false, "should_attach": true/false, "reason": "中文理由", "selected_paths": ["绝对路径1", "绝对路径2"]}\n'
        "要求：\n"
        "1. 最多选择 3 个文件。\n"
        "2. 只能从给定列表里选。\n"
        "3. 如果邮件只是讨论、确认、询问，但问题涉及目录下材料内容，可以 use_catalog=true。\n"
        "4. 只有高置信度时才返回 should_attach=true。"
    )
    response = client.messages.create(
        model=model,
        max_tokens=520,
        messages=[{"role": "user", "content": prompt}],
    )
    text = "\n".join(getattr(block, "text", "") for block in response.content).strip()
    match = re.search(r"\{.*\}", text, re.S)
    if not match:
        return {
            "use_catalog": False,
            "should_attach": False,
            "reason": "Attachment selector returned non-JSON output.",
            "selected_paths": [],
        }
    try:
        payload = json.loads(match.group(0))
    except Exception:
        return {
            "use_catalog": False,
            "should_attach": False,
            "reason": "Attachment selector returned invalid JSON.",
            "selected_paths": [],
        }
    selected_paths = []
    allowed_paths = {item["path"] for item in candidate_files}
    for value in payload.get("selected_paths", []) or []:
        normalized = str(value).strip()
        if normalized in allowed_paths and normalized not in selected_paths:
            selected_paths.append(normalized)
    return {
        "use_catalog": bool(payload.get("use_catalog")) and bool(selected_paths),
        "should_attach": bool(payload.get("should_attach")) and bool(selected_paths),
        "reason": str(payload.get("reason", "")).strip(),
        "selected_paths": selected_paths[:3],
    }


def _format_catalog_reference_context(candidate_files: list[dict], selected_paths: list[str]) -> str:
    if not selected_paths:
        return ""
    items_by_path = {str(item.get("path", "")).strip(): item for item in candidate_files}
    sections: list[str] = []
    for idx, path_value in enumerate(selected_paths, start=1):
        item = items_by_path.get(path_value)
        if not item:
            continue
        sections.extend(
            [
                f"[资料 {idx}] 文件名：{item.get('name', '')}",
                f"路径：{item.get('path', '')}",
                f"摘要：{item.get('summary', '')}",
                f"标签：{item.get('tags', '')}",
                f"适用场景：{item.get('use_cases', '')}",
                f"关键词：{item.get('keywords', '')}",
                f"内容摘录：{_truncate_text(str(item.get('excerpt', '')), limit=1200)}",
                "",
            ]
        )
    return "\n".join(sections).strip()


def _attach_local_files_to_message(mime: EmailMessage, attachments: list[Path]) -> None:
    for path in attachments:
        data = path.read_bytes()
        content_type, encoding = mimetypes.guess_type(str(path))
        if encoding:
            content_type = "application/octet-stream"
        maintype, subtype = (content_type or "application/octet-stream").split("/", 1)
        mime.add_attachment(data, maintype=maintype, subtype=subtype, filename=path.name)


def send_gmail_reply(subject: str, reply_body: str, unread_only: bool = False, local_attachment_paths: Optional[list[str]] = None) -> dict:
    _require_compose_scope()
    message = find_gmail_message_by_subject(subject, unread_only=unread_only)
    reply_target = message.get("reply_to") or message.get("from", "")
    _, reply_email = parseaddr(reply_target)
    if not reply_email:
        raise RuntimeError(f"无法从邮件发件人里解析回复地址：{reply_target}")

    original_subject = message.get("subject", subject).strip()
    draft_subject = original_subject if original_subject.lower().startswith("re:") else f"Re: {original_subject}"
    my_email = _get_authenticated_email()
    original_to = _normalize_email_list(message.get("to", ""))
    original_cc = _normalize_email_list(message.get("cc", ""))
    primary_to = reply_email.strip().lower()
    cc_list = [
        email
        for email in (original_to + original_cc)
        if email and email not in {my_email, primary_to}
    ]

    mime = EmailMessage()
    mime["To"] = reply_email
    if cc_list:
        mime["Cc"] = ", ".join(cc_list)
    mime["Subject"] = draft_subject
    if message.get("message_id_header"):
        mime["In-Reply-To"] = message["message_id_header"]
        references = message.get("references", "").strip()
        mime["References"] = (references + " " + message["message_id_header"]).strip() if references else message["message_id_header"]
    mime.set_content(_sanitize_reply_body(reply_body) + _format_quoted_reply_block(message))
    selected_attachment_paths: list[str] = []
    attachment_objects: list[Path] = []
    for raw_path in local_attachment_paths or []:
        candidate = Path(raw_path).expanduser().resolve()
        if candidate.exists() and candidate.is_file():
            attachment_objects.append(candidate)
            selected_attachment_paths.append(str(candidate))
    if attachment_objects:
        _attach_local_files_to_message(mime, attachment_objects)

    raw = base64.urlsafe_b64encode(mime.as_bytes()).decode("utf-8")
    payload = {
        "threadId": message.get("threadId", ""),
        "raw": raw,
    }
    sent = _gmail_api_post("users/me/messages/send", payload)
    return {
        "message_id": sent.get("id", ""),
        "thread_id": sent.get("threadId", "") or message.get("threadId", ""),
        "to": reply_email,
        "cc": cc_list,
        "subject": draft_subject,
        "matched_subject": original_subject,
        "local_attachment_paths": selected_attachment_paths,
    }


def mark_gmail_message_read(message_id: str) -> dict:
    _require_modify_scope()
    return _gmail_api_post(
        f"users/me/messages/{message_id}/modify",
        {"removeLabelIds": ["UNREAD"]},
    )


def _append_markdown(log_path: Path, email_data: dict, assessment: Optional[dict]) -> None:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(f"## {_now()}\n")
        handle.write(f"- 发件人: {email_data.get('from', '未知')}\n")
        handle.write(f"- 主题: {email_data.get('subject', '无')}\n")
        handle.write(f"- 时间: {email_data.get('date', '未知')}\n")
        handle.write(f"- 预览: {email_data.get('snippet', '无')}\n")
        handle.write("- 正文:\n\n```\n")
        handle.write((email_data.get("body") or "").strip() + "\n")
        handle.write("```\n\n")
        if assessment:
            handle.write(f"- 建议回复: {assessment.get('need_reply', 'unknown')}\n")
            handle.write(f"- 判断理由: {assessment.get('reason', '无')}\n")
            local_attachment_source = (assessment.get("local_attachment_source") or "").strip()
            if local_attachment_source:
                handle.write(f"- 本地附件来源: {local_attachment_source}\n")
            local_attachment_paths = assessment.get("local_attachment_paths") or []
            local_attachment_reason = (assessment.get("local_attachment_reason") or "").strip()
            if local_attachment_paths:
                handle.write("- 本地回复附件:\n")
                for path in local_attachment_paths:
                    handle.write(f"  - {path}\n")
            elif local_attachment_reason:
                handle.write(f"- 本地回复附件: 无 ({local_attachment_reason})\n")
            if assessment.get("sent_reply_id"):
                handle.write(f"- 已发送回复: yes\n")
                handle.write(f"- 回复主题: {assessment.get('sent_reply_subject', '无')}\n")
                handle.write(f"- 回复消息 ID: {assessment.get('sent_reply_id', '无')}\n")
            elif assessment.get("sent_reply_error"):
                handle.write(f"- 已发送回复: no\n")
                handle.write(f"- 发送错误: {assessment.get('sent_reply_error', '无')}\n")
            if assessment.get("marked_read"):
                handle.write("- 已标记为已读: yes\n")
            elif assessment.get("mark_read_error"):
                handle.write(f"- 已标记为已读: no\n")
                handle.write(f"- 标记已读错误: {assessment.get('mark_read_error', '无')}\n")
            draft = (assessment.get("draft_reply") or "").strip()
            if draft:
                handle.write("- 建议回复草稿:\n\n```\n")
                handle.write(draft + "\n")
                handle.write("```\n\n")


def _recent_entries(log_path: Path, limit: int = 5) -> list[dict]:
    if not log_path.exists():
        return []
    text = log_path.read_text(encoding="utf-8", errors="replace")
    chunks = [chunk.strip() for chunk in text.split("## ") if chunk.strip()]
    entries: list[dict] = []
    for chunk in reversed(chunks):
        lines = chunk.splitlines()
        if not lines:
            continue
        entry = {
            "timestamp": lines[0].strip(),
            "from": "",
            "subject": "",
            "date": "",
            "snippet": "",
            "body": "",
            "need_reply": "",
            "reason": "",
            "draft_reply": "",
        }
        body_lines: list[str] = []
        draft_lines: list[str] = []
        in_body = False
        in_draft = False
        in_fence = False
        for line in lines[1:]:
            stripped = line.strip()
            if stripped == "```":
                in_fence = not in_fence
                continue
            if stripped.startswith("- 发件人:"):
                entry["from"] = stripped.split(":", 1)[1].strip()
                in_body = in_draft = False
            elif stripped.startswith("- 主题:"):
                entry["subject"] = stripped.split(":", 1)[1].strip()
                in_body = in_draft = False
            elif stripped.startswith("- 时间:"):
                entry["date"] = stripped.split(":", 1)[1].strip()
                in_body = in_draft = False
            elif stripped.startswith("- 预览:"):
                entry["snippet"] = stripped.split(":", 1)[1].strip()
                in_body = in_draft = False
            elif stripped.startswith("- 正文:"):
                in_body, in_draft = True, False
            elif stripped.startswith("- 建议回复:"):
                entry["need_reply"] = stripped.split(":", 1)[1].strip()
                in_body = in_draft = False
            elif stripped.startswith("- 判断理由:"):
                entry["reason"] = stripped.split(":", 1)[1].strip()
                in_body = in_draft = False
            elif stripped.startswith("- 建议回复草稿:"):
                in_body, in_draft = False, True
            elif in_fence and in_body:
                body_lines.append(line)
            elif in_fence and in_draft:
                draft_lines.append(line)
        entry["body"] = "\n".join(body_lines).strip()
        entry["draft_reply"] = "\n".join(draft_lines).strip()
        entries.append(entry)
        if len(entries) >= limit:
            break
    return entries


def _is_no_reply_message(email_data: dict) -> bool:
    subject = str(email_data.get("subject", "")).strip()
    sender = str(email_data.get("from", "")).strip()
    thread_markers = [
        "[Yuan-lab-LLM/ClawManager]",
    ]
    return any(marker in subject or marker in sender for marker in thread_markers)


def _safe_fetch_current_unread_entries(limit: int = 5) -> list[dict]:
    try:
        query = os.getenv(
            "GMAIL_WATCHDOG_QUERY",
            "is:unread -in:drafts -in:sent -category:promotions -category:social",
        )
        payload = _gmail_api_get(
            "users/me/messages",
            {"q": query, "maxResults": max(1, int(limit))},
        )
        messages = payload.get("messages", []) or []
        entries: list[dict] = []
        for item in messages[:limit]:
            message_id = item.get("id")
            if not message_id:
                continue
            detail_payload = _gmail_api_get(
                f"users/me/messages/{message_id}",
                {"format": "full"},
            )
            headers = _message_headers(detail_payload.get("payload", {}))
            body = _extract_plain_text(detail_payload.get("payload", {})) or _clean_text(
                detail_payload.get("snippet", "")
            )
            date_value = headers.get("date", "")
            try:
                date_text = (
                    parsedate_to_datetime(date_value).astimezone().strftime("%Y-%m-%d %H:%M:%S")
                    if date_value
                    else ""
                )
            except Exception:
                date_text = date_value
            entries.append(
                {
                    "id": detail_payload.get("id", message_id),
                    "threadId": detail_payload.get("threadId", ""),
                    "from": headers.get("from", ""),
                    "subject": headers.get("subject", ""),
                    "date": date_text,
                    "snippet": _clean_text(detail_payload.get("snippet", "")),
                    "body": body,
                }
            )
        return entries
    except Exception:
        return []


@dataclass
class GmailWatchdogSnapshot:
    running: bool = False
    interval_seconds: int = 60
    log_path: str = str(DEFAULT_LOG_PATH)
    last_scan_at: str = ""
    last_detection_at: str = ""
    last_error: str = ""
    last_message_from: str = ""
    last_subject: str = ""
    last_reply_assessment: str = ""
    last_snippet: str = ""
    last_body: str = ""
    last_need_reply: str = ""
    last_reply_reason: str = ""
    last_reply_draft: str = ""
    last_sent_reply_id: str = ""
    last_sent_reply_subject: str = ""
    last_marked_read: bool = False
    current_unread_entries: list[dict] | None = None
    latest_unread: dict | None = None
    detections: int = 0
    token_path: str = str(DEFAULT_TOKEN_PATH)


@dataclass
class GmailAttachmentCatalogBuildSnapshot:
    running: bool = False
    ready: bool = False
    catalog_path: str = str(DEFAULT_ATTACHMENT_CATALOG_PATH)
    configured_dirs: list[str] | None = None
    total_files: int = 0
    processed_files: int = 0
    indexed_files: int = 0
    skipped_files: int = 0
    current_file: str = ""
    last_started_at: str = ""
    last_finished_at: str = ""
    last_error: str = ""
    last_message: str = ""


class GmailAttachmentCatalogBuilder:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._thread: Optional[threading.Thread] = None
        self._state_path = DEFAULT_CATALOG_BUILD_STATE_PATH
        self._snapshot = GmailAttachmentCatalogBuildSnapshot(
            catalog_path=str(_attachment_catalog_path()),
            configured_dirs=[str(path) for path in _gmail_local_attachment_roots()],
        )
        self._load_state()

    def _load_state(self) -> None:
        if not self._state_path.exists():
            return
        try:
            payload = _load_json(self._state_path)
        except Exception:
            return
        configured_dirs = payload.get("configured_dirs") or []
        self._snapshot = GmailAttachmentCatalogBuildSnapshot(
            running=False,
            ready=bool(payload.get("ready")),
            catalog_path=str(payload.get("catalog_path") or _attachment_catalog_path()),
            configured_dirs=[str(item) for item in configured_dirs if str(item).strip()],
            total_files=int(payload.get("total_files") or 0),
            processed_files=int(payload.get("processed_files") or 0),
            indexed_files=int(payload.get("indexed_files") or 0),
            skipped_files=int(payload.get("skipped_files") or 0),
            current_file="",
            last_started_at=str(payload.get("last_started_at") or ""),
            last_finished_at=str(payload.get("last_finished_at") or ""),
            last_error=str(payload.get("last_error") or ""),
            last_message=str(payload.get("last_message") or ""),
        )

    def _save_state_unlocked(self) -> None:
        _write_json(
            self._state_path,
            {
                "ready": self._snapshot.ready,
                "catalog_path": self._snapshot.catalog_path,
                "configured_dirs": list(self._snapshot.configured_dirs or []),
                "total_files": self._snapshot.total_files,
                "processed_files": self._snapshot.processed_files,
                "indexed_files": self._snapshot.indexed_files,
                "skipped_files": self._snapshot.skipped_files,
                "last_started_at": self._snapshot.last_started_at,
                "last_finished_at": self._snapshot.last_finished_at,
                "last_error": self._snapshot.last_error,
                "last_message": self._snapshot.last_message,
            },
        )

    def _snapshot_dict_unlocked(self) -> dict:
        configured_dirs = [str(path) for path in _gmail_local_attachment_roots()]
        ready = (
            self._snapshot.ready
            and not self._snapshot.running
            and configured_dirs == list(self._snapshot.configured_dirs or [])
            and self._snapshot.catalog_path == str(_attachment_catalog_path())
            and Path(self._snapshot.catalog_path).exists()
        )
        return {
            "running": self._snapshot.running,
            "ready": ready,
            "catalog_path": str(_attachment_catalog_path()),
            "configured_dirs": configured_dirs,
            "total_files": self._snapshot.total_files,
            "processed_files": self._snapshot.processed_files,
            "indexed_files": self._snapshot.indexed_files,
            "skipped_files": self._snapshot.skipped_files,
            "current_file": self._snapshot.current_file,
            "last_started_at": self._snapshot.last_started_at,
            "last_finished_at": self._snapshot.last_finished_at,
            "last_error": self._snapshot.last_error,
            "last_message": self._snapshot.last_message,
        }

    def status(self) -> dict:
        with self._lock:
            return self._snapshot_dict_unlocked()

    def start(self) -> dict:
        roots = _gmail_local_attachment_roots()
        if not roots:
            raise RuntimeError("请先配置 Gmail 附件白名单文件夹，再初始化附件知识库。")
        with self._lock:
            if self._thread and self._thread.is_alive():
                return self._snapshot_dict_unlocked()
            self._snapshot = GmailAttachmentCatalogBuildSnapshot(
                running=True,
                ready=False,
                catalog_path=str(_attachment_catalog_path()),
                configured_dirs=[str(path) for path in roots],
                total_files=0,
                processed_files=0,
                indexed_files=0,
                skipped_files=0,
                current_file="",
                last_started_at=_now(),
                last_finished_at="",
                last_error="",
                last_message="正在分析附件白名单文件夹里的文档内容...",
            )
            self._save_state_unlocked()
            self._thread = threading.Thread(target=self._build_loop, daemon=True)
            self._thread.start()
            return self._snapshot_dict_unlocked()

    def _build_loop(self) -> None:
        roots = _gmail_local_attachment_roots()
        catalog_path = _attachment_catalog_path()
        supported = _supported_catalog_suffixes()
        all_files: list[Path] = []
        for root in roots:
            try:
                for candidate in root.rglob("*"):
                    if not candidate.is_file():
                        continue
                    if candidate.suffix.lower() not in supported:
                        continue
                    all_files.append(candidate.resolve())
            except Exception:
                continue
        deduped: list[Path] = []
        seen: set[Path] = set()
        for item in sorted(all_files, key=lambda path: str(path).lower()):
            if item not in seen:
                deduped.append(item)
                seen.add(item)
        with self._lock:
            self._snapshot.total_files = len(deduped)
            self._snapshot.current_file = ""
            self._save_state_unlocked()

        entries: list[dict] = []
        try:
            for index, file_path in enumerate(deduped, start=1):
                with self._lock:
                    self._snapshot.current_file = str(file_path)
                    self._snapshot.processed_files = index - 1
                    self._snapshot.last_message = f"正在分析第 {index}/{len(deduped)} 个文件：{file_path.name}"
                    self._save_state_unlocked()
                try:
                    text = _extract_local_file_text(file_path)
                    if not _clean_text(text):
                        raise RuntimeError("未提取到可用文本")
                    entries.append(_build_catalog_entry(file_path, text))
                    with self._lock:
                        self._snapshot.indexed_files += 1
                except Exception:
                    with self._lock:
                        self._snapshot.skipped_files += 1
                with self._lock:
                    self._snapshot.processed_files = index
                    self._save_state_unlocked()
            catalog_path.parent.mkdir(parents=True, exist_ok=True)
            catalog_path.write_text(_render_attachment_catalog(entries), encoding="utf-8")
            with self._lock:
                self._snapshot.running = False
                self._snapshot.ready = True
                self._snapshot.current_file = ""
                self._snapshot.last_finished_at = _now()
                self._snapshot.last_message = (
                    f"初始化完成：共扫描 {len(deduped)} 个文件，成功建立 {len(entries)} 条知识库记录。"
                )
                self._save_state_unlocked()
        except Exception as exc:
            with self._lock:
                self._snapshot.running = False
                self._snapshot.ready = False
                self._snapshot.current_file = ""
                self._snapshot.last_error = str(exc)
                self._snapshot.last_finished_at = _now()
                self._snapshot.last_message = f"初始化失败：{exc}"
                self._save_state_unlocked()


CATALOG_BUILDER = GmailAttachmentCatalogBuilder()


class GmailWatchdog:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._stop_event = threading.Event()
        self._thread: Optional[threading.Thread] = None
        self._snapshot = GmailWatchdogSnapshot()
        self._state_path = DEFAULT_STATE_PATH
        self._seen_ids: dict[str, str] = {}

    def _snapshot_dict_unlocked(self) -> dict:
        current_unread_entries = list(self._snapshot.current_unread_entries or [])
        latest_unread = dict(self._snapshot.latest_unread or {})
        return {
            "running": self._snapshot.running,
            "interval_seconds": self._snapshot.interval_seconds,
            "log_path": self._snapshot.log_path,
            "last_scan_at": self._snapshot.last_scan_at,
            "last_detection_at": self._snapshot.last_detection_at,
            "last_error": self._snapshot.last_error,
            "last_message_from": self._snapshot.last_message_from,
            "last_subject": self._snapshot.last_subject,
            "last_reply_assessment": self._snapshot.last_reply_assessment,
            "last_snippet": self._snapshot.last_snippet,
            "last_body": self._snapshot.last_body,
            "last_need_reply": self._snapshot.last_need_reply,
            "last_reply_reason": self._snapshot.last_reply_reason,
            "last_reply_draft": self._snapshot.last_reply_draft,
            "last_sent_reply_id": self._snapshot.last_sent_reply_id,
            "last_sent_reply_subject": self._snapshot.last_sent_reply_subject,
            "last_marked_read": self._snapshot.last_marked_read,
            "detections": self._snapshot.detections,
            "token_path": self._snapshot.token_path,
            "current_unread_entries": current_unread_entries,
            "latest_unread": latest_unread,
            "recent_entries": _recent_entries(Path(self._snapshot.log_path), limit=5),
        }

    def _load_state(self) -> None:
        if not self._state_path.exists():
            self._seen_ids = {}
            return
        try:
            payload = _load_json(self._state_path)
            self._seen_ids = payload.get("seen_ids", {})
        except Exception:
            self._seen_ids = {}

    def _save_state(self) -> None:
        _write_json(self._state_path, {"seen_ids": self._seen_ids})

    def start(self, interval_seconds: int = 60, log_path: Optional[str] = None) -> dict:
        with self._lock:
            if self._thread and self._thread.is_alive():
                return self._snapshot_dict_unlocked()
            self._stop_event.clear()
            self._snapshot = GmailWatchdogSnapshot(
                running=True,
                interval_seconds=max(30, int(interval_seconds)),
                log_path=str(Path(log_path).expanduser().resolve()) if log_path else str(DEFAULT_LOG_PATH),
                token_path=str(Path(os.getenv("GMAIL_TOKEN_FILE", str(DEFAULT_TOKEN_PATH))).expanduser().resolve()),
            )
            self._load_state()
            self._thread = threading.Thread(target=self._loop, daemon=True)
            self._thread.start()
            return self._snapshot_dict_unlocked()

    def stop(self) -> dict:
        with self._lock:
            self._stop_event.set()
            self._snapshot.running = False
            return self._snapshot_dict_unlocked()

    def status(self) -> dict:
        with self._lock:
            return self._snapshot_dict_unlocked()

    def _assess_reply_need(
        self,
        sender: str,
        subject: str,
        content: str,
        has_attachments: bool = False,
        attachment_context_available: bool = False,
    ) -> dict:
        api_key = os.getenv("KIMI_API_KEY") or os.getenv("ANTHROPIC_API_KEY")
        if not api_key:
            return {
                "need_reply": "unknown",
                "reason": "No Kimi/Anthropic API key available for reply assessment.",
                "draft_reply": "",
            }
        model = os.getenv("KIMI_MODEL", os.getenv("ANTHROPIC_MODEL", "kimi-k2.5"))
        base_url = os.getenv("KIMI_BASE_URL", os.getenv("ANTHROPIC_BASE_URL", "https://api.moonshot.cn/anthropic"))
        client = anthropic.Anthropic(api_key=api_key, base_url=base_url)
        attachment_rule = ""
        if has_attachments and attachment_context_available:
            attachment_rule = (
                "这封邮件带有附件，且系统已经成功提取了附件内容。\n"
                "如果邮件请求的答复依赖附件内容，你必须基于附件内容直接形成最终回复，"
                "不要写“我会先查看附件/稍后回复/尽快整理后反馈”这类过渡性表述。\n"
                "如果附件内容已经足够支持答复，就直接给出结论、摘要、分析或问题答复。\n\n"
            )
        elif has_attachments and not attachment_context_available:
            attachment_rule = (
                "这封邮件带有附件，但系统当前没有成功提取附件正文。\n"
                "如果邮件明确要求基于附件内容答复，你可以说明当前无法读取附件内容，并请对方提供可读文本或关键数据。\n\n"
            )

        prompt = (
            "你在做 Gmail 邮件值守。请判断下面这封未读邮件是否需要回复。\n\n"
            f"发件人：{sender}\n"
            f"主题：{subject}\n"
            f"正文：{content}\n\n"
            f"{attachment_rule}"
            "请只用下面格式回答：\n"
            "NEED_REPLY: yes 或 no\n"
            "REASON: 一句简短中文理由\n"
            "DRAFT_REPLY: 如果需要回复，给出可以直接发送的完整中文回复；如果不需要，写 N/A"
        )
        response = client.messages.create(
            model=model,
            max_tokens=520,
            messages=[{"role": "user", "content": prompt}],
        )
        text = "\n".join(getattr(block, "text", "") for block in response.content).strip()
        need_reply = "unknown"
        reason = ""
        draft_reply = ""
        for line in text.splitlines():
            if line.startswith("NEED_REPLY:"):
                value = line.split(":", 1)[1].strip().lower()
                need_reply = "yes" if value.startswith("y") else "no" if value.startswith("n") else "unknown"
            elif line.startswith("REASON:"):
                reason = line.split(":", 1)[1].strip()
            elif line.startswith("DRAFT_REPLY:"):
                draft_reply = line.split(":", 1)[1].strip()
        return {
            "need_reply": need_reply,
            "reason": reason or "No reason returned.",
            "draft_reply": "" if draft_reply == "N/A" else draft_reply,
        }

    def _generate_final_reply(
        self,
        sender: str,
        subject: str,
        content: str,
        has_attachments: bool = False,
        attachment_context_available: bool = False,
    ) -> str:
        api_key = os.getenv("KIMI_API_KEY") or os.getenv("ANTHROPIC_API_KEY")
        if not api_key:
            return ""
        model = os.getenv("KIMI_MODEL", os.getenv("ANTHROPIC_MODEL", "kimi-k2.5"))
        base_url = os.getenv("KIMI_BASE_URL", os.getenv("ANTHROPIC_BASE_URL", "https://api.moonshot.cn/anthropic"))
        client = anthropic.Anthropic(api_key=api_key, base_url=base_url)

        attachment_rule = ""
        if has_attachments and attachment_context_available:
            attachment_rule = (
                "这封邮件带有附件，且系统已经成功提取了附件内容。\n"
                "你必须基于邮件正文和附件内容直接形成一封可立即发送的最终回复。\n"
                "不要写“我会先查看附件/稍后回复/尽快整理后反馈”这类过渡性表述。\n\n"
            )
        elif has_attachments and not attachment_context_available:
            attachment_rule = (
                "这封邮件带有附件，但系统当前没有成功提取附件正文。\n"
                "如果邮件明确要求基于附件内容答复，请如实说明当前无法读取附件，并请对方提供可读文本或关键数据。\n\n"
            )

        prompt = (
            "你在做 Gmail 自动回复。\n"
            "请直接输出一封可以发送的、自然专业的中文回复正文，不要加解释、不要加标题、不要加格式标签。\n"
            "不要使用 Markdown，不要使用星号加粗，不要使用项目符号列表，直接输出普通邮件正文。\n\n"
            f"发件人：{sender}\n"
            f"主题：{subject}\n"
            f"正文与附件信息：{content}\n\n"
            f"{attachment_rule}"
            "要求：\n"
            "1. 直接给出最终回复，不要输出分析过程。\n"
            "2. 如果邮件请求基于附件内容答复，就必须体现附件中的关键结论。\n"
            "3. 语言简洁、自然、像真实商务邮件回复。\n"
            "4. 分段即可，不要用 *、**、- 这类 Markdown 标记。"
        )
        response = client.messages.create(
            model=model,
            max_tokens=520,
            messages=[{"role": "user", "content": prompt}],
        )
        return _sanitize_reply_body("\n".join(getattr(block, "text", "") for block in response.content).strip())

    def _record_detection(self, email_data: dict, assessment: dict) -> None:
        _append_markdown(Path(self._snapshot.log_path), email_data, assessment)
        self._snapshot.last_detection_at = _now()
        self._snapshot.last_message_from = email_data.get("from", "")
        self._snapshot.last_subject = email_data.get("subject", "")
        self._snapshot.last_snippet = email_data.get("snippet", "")
        self._snapshot.last_body = (email_data.get("body", "") or "").strip()
        self._snapshot.last_need_reply = assessment.get("need_reply", "unknown")
        self._snapshot.last_reply_reason = assessment.get("reason", "")
        self._snapshot.last_reply_draft = (assessment.get("draft_reply", "") or "").strip()
        self._snapshot.last_sent_reply_id = assessment.get("sent_reply_id", "")
        self._snapshot.last_sent_reply_subject = assessment.get("sent_reply_subject", "")
        self._snapshot.last_marked_read = bool(assessment.get("marked_read"))
        self._snapshot.last_reply_assessment = (
            f"{email_data.get('subject', '')}: need_reply={assessment.get('need_reply', 'unknown')}; "
            f"reason={assessment.get('reason', '')}"
        )
        if assessment.get("local_attachment_paths"):
            self._snapshot.last_reply_assessment += (
                f"; local_attachments={len(assessment.get('local_attachment_paths') or [])}"
            )
        self._snapshot.detections += 1

    def _fetch_unread_messages(self) -> list[dict]:
        query = os.getenv(
            "GMAIL_WATCHDOG_QUERY",
            "is:unread -in:drafts -in:sent -category:promotions -category:social",
        )
        payload = _gmail_api_get(
            "users/me/messages",
            {"q": query, "maxResults": 10},
        )
        return payload.get("messages", [])

    def _fetch_message_detail(self, message_id: str) -> dict:
        return _fetch_message_detail_by_id(message_id)

    def _loop(self) -> None:
        while not self._stop_event.is_set():
            try:
                unread = self._fetch_unread_messages()
                self._snapshot.last_scan_at = _now()
                self._snapshot.last_error = ""
                current_entries: list[dict] = []
                for item in unread[:5]:
                    message_id = item.get("id")
                    if not message_id:
                        continue
                    try:
                        current_entries.append(self._fetch_message_detail(message_id))
                    except Exception:
                        continue
                self._snapshot.current_unread_entries = current_entries
                self._snapshot.latest_unread = current_entries[0] if current_entries else {}
                for item in unread:
                    message_id = item.get("id")
                    if not message_id or message_id in self._seen_ids:
                        continue
                    email_data = self._fetch_message_detail(message_id)
                    if _is_no_reply_message(email_data):
                        assessment = {
                            "need_reply": "no",
                            "reason": "命中不回复规则：来自 [Yuan-lab-LLM/ClawManager] 的邮件不自动回复。",
                            "draft_reply": "",
                            "local_attachment_reason": "",
                            "local_attachment_paths": [],
                            "local_attachment_source": "",
                        }
                        try:
                            mark_gmail_message_read(message_id)
                            assessment["marked_read"] = True
                        except Exception as mark_exc:
                            assessment["mark_read_error"] = str(mark_exc)
                        self._seen_ids[message_id] = email_data.get("threadId", "") or "seen"
                        self._save_state()
                        self._record_detection(email_data, assessment)
                        continue
                    content_for_assessment = (email_data.get("body", "") or email_data.get("snippet", "")).strip()
                    attachment_context = (email_data.get("attachment_context") or "").strip()
                    if attachment_context:
                        content_for_assessment = (
                            f"{content_for_assessment}\n\n附件内容如下：\n{attachment_context}"
                        ).strip()
                    candidate_files = _catalog_attachment_candidates(limit=80)
                    if not candidate_files:
                        candidate_files = _list_local_attachment_candidates(limit=80)
                    material_plan = {
                        "use_catalog": False,
                        "should_attach": False,
                        "reason": "",
                        "selected_paths": [],
                    }
                    try:
                        material_plan = _select_local_reply_attachments(
                            email_data.get("from", ""),
                            email_data.get("subject", ""),
                            content_for_assessment,
                            candidate_files,
                        )
                    except Exception as material_exc:
                        material_plan["reason"] = f"本地资料匹配失败：{material_exc}"
                    selected_paths = list(material_plan.get("selected_paths") or [])
                    if material_plan.get("use_catalog") and selected_paths:
                        catalog_context = _format_catalog_reference_context(candidate_files, selected_paths)
                        if catalog_context:
                            content_for_assessment = (
                                f"{content_for_assessment}\n\n【本地资料知识库相关文档】\n{catalog_context}"
                            ).strip()
                    assessment = self._assess_reply_need(
                        email_data.get("from", ""),
                        email_data.get("subject", ""),
                        content_for_assessment,
                        has_attachments=bool(email_data.get("attachments")),
                        attachment_context_available=bool(attachment_context),
                    )
                    if assessment.get("need_reply") == "yes" and not (assessment.get("draft_reply") or "").strip():
                        try:
                            assessment["draft_reply"] = self._generate_final_reply(
                                email_data.get("from", ""),
                                email_data.get("subject", ""),
                                content_for_assessment,
                                has_attachments=bool(email_data.get("attachments")),
                                attachment_context_available=bool(attachment_context),
                            )
                        except Exception as reply_exc:
                            assessment["draft_reply_error"] = str(reply_exc)
                    assessment["local_attachment_reason"] = ""
                    assessment["local_attachment_paths"] = []
                    assessment["local_attachment_source"] = ""
                    if material_plan.get("reason"):
                        assessment["local_attachment_reason"] = material_plan.get("reason", "")
                    if material_plan.get("use_catalog") and selected_paths:
                        assessment["local_attachment_source"] = "catalog"
                    if material_plan.get("should_attach") and selected_paths:
                        assessment["local_attachment_paths"] = selected_paths
                    if assessment.get("need_reply") == "yes" and (assessment.get("draft_reply") or "").strip():
                        try:
                            sent = send_gmail_reply(
                                email_data.get("subject", ""),
                                (assessment.get("draft_reply") or "").strip(),
                                local_attachment_paths=list(assessment.get("local_attachment_paths") or []),
                            )
                            assessment["sent_reply_id"] = sent.get("message_id", "")
                            assessment["sent_reply_subject"] = sent.get("subject", "")
                            assessment["local_attachment_paths"] = sent.get("local_attachment_paths", assessment.get("local_attachment_paths") or [])
                        except Exception as send_exc:
                            assessment["sent_reply_error"] = str(send_exc)
                    try:
                        mark_gmail_message_read(message_id)
                        assessment["marked_read"] = True
                    except Exception as mark_exc:
                        assessment["mark_read_error"] = str(mark_exc)
                    self._seen_ids[message_id] = email_data.get("threadId", "") or "seen"
                    self._save_state()
                    self._record_detection(email_data, assessment)
            except Exception as exc:
                self._snapshot.last_scan_at = _now()
                self._snapshot.last_error = str(exc)
            self._stop_event.wait(self._snapshot.interval_seconds)


WATCHDOG = GmailWatchdog()


def start_gmail_watchdog(interval_seconds: int = 60, log_path: Optional[str] = None) -> dict:
    return WATCHDOG.start(interval_seconds=interval_seconds, log_path=log_path)


def stop_gmail_watchdog() -> dict:
    return WATCHDOG.stop()


def get_gmail_watchdog_status() -> dict:
    return WATCHDOG.status()


def reset_gmail_watchdog_seen() -> dict:
    WATCHDOG._seen_ids = {}
    WATCHDOG._save_state()
    return WATCHDOG.status()


def start_gmail_attachment_catalog_build() -> dict:
    return CATALOG_BUILDER.start()


def get_gmail_attachment_catalog_build_status() -> dict:
    return CATALOG_BUILDER.status()
